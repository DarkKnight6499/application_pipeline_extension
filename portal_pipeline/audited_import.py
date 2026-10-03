"""Read one explicitly selected application and run the existing audit in place."""
from __future__ import annotations

import hashlib
import json
import os
import re
import subprocess
import sys
from pathlib import Path
from urllib.parse import urlsplit


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def posting_url(value):
    url = str(value or "").strip()
    parsed = urlsplit(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:
        raise ValueError("Provide a valid posting URL without embedded credentials.")
    return url


def tracker_context(source, application_id):
    from openpyxl import load_workbook
    workbook = load_workbook(source / "Applications.xlsx", read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        headers = next(rows)
        required = {"Application ID", "Company", "Role Title", "Link", "Status"}
        if not required.issubset(headers):
            raise ValueError("The tracker is missing required application columns.")
        index = headers.index("Application ID")
        matches = [dict(zip(headers, row)) for row in rows if row[index] == application_id]
    finally:
        workbook.close()
    if len(matches) != 1:
        raise ValueError("The application marker must match exactly one tracker row.")
    row = matches[0]
    if row["Status"] not in {"New", "To Apply"}:
        raise ValueError("Only New or To Apply applications can be imported. This prototype never changes tracker status.")
    url = posting_url(row["Link"])
    if not row["Company"] or not row["Role Title"]:
        raise ValueError("The tracker row needs a company and role title.")
    return {"application_id": application_id, "company": str(row["Company"]), "role": str(row["Role Title"]),
            "url": url, "tracker_status": row["Status"]}


def application_files(source, folder):
    applications = (source / "Applications").resolve()
    folder = Path(folder).resolve()
    if not folder.is_relative_to(applications) or folder == applications or not folder.is_dir():
        raise ValueError("Select one existing folder inside the source Applications directory.")
    inputs = folder / "_inputs"
    # Reject duplicate working inputs because the existing audit prefers top-level files.
    if (folder / "resume_content.json").exists() or list(folder.glob("Keywords_*.json")):
        raise ValueError("Finish the current workflow and archive its working inputs before import.")
    keywords = list(inputs.glob("Keywords_*.json"))
    jds = list(folder.glob("JD_*.docx"))
    if len(keywords) != 1 or len(jds) != 1:
        raise ValueError("Import requires exactly one archived Keywords file and one JD document.")
    paths = [folder / ".application_id", folder / "Yazad_Madan.docx", inputs / "resume_content.json", keywords[0], jds[0]]
    paths += list(folder.glob("CoverLetter_*.docx"))
    for path in paths:
        if not path.is_file() or not path.resolve().is_relative_to(folder):
            raise ValueError("Required application files must exist inside the selected folder.")
    return folder, paths, keywords[0]


def fingerprint(folder, paths):
    return {str(path.relative_to(folder)): digest(path) for path in paths}


def inspect_application(source, folder):
    folder, paths, keywords_path = application_files(source, folder)
    marker = (folder / ".application_id").read_text(encoding="utf-8").strip()
    if not re.fullmatch(r"[1-9][0-9]*", marker):
        raise ValueError("The selected folder needs a valid positive Application ID marker.")
    context = tracker_context(source, int(marker))
    snapshot = fingerprint(folder, paths)
    content = json.loads((folder / "_inputs/resume_content.json").read_text(encoding="utf-8"))
    keywords = json.loads(keywords_path.read_text(encoding="utf-8"))
    normalized = lambda value: re.sub(r"[^a-z0-9]", "", str(value).lower())
    if normalized(keywords.get("company")) != normalized(context["company"]) or normalized(keywords.get("role")) != normalized(context["role"]):
        raise ValueError("Archived JD requirements and the exact tracker row identify different applications.")
    if not isinstance(content, dict) or keywords.get("requirements_reviewed") is not True:
        raise ValueError("Import requires archived content and a completed JD requirements review.")
    result = subprocess.run([sys.executable, "-B", str(source / "_Reference/audit_application.py"), str(folder)],
                            cwd=source, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=45,
                            env={**os.environ, "PYTHONIOENCODING": "utf-8", "PYTHONDONTWRITEBYTECODE": "1"})
    report = result.stdout + result.stderr
    if result.returncode:
        raise ValueError("The current application audit failed. Fix it through the existing workflow before import.\n" + report[-6000:])
    check_source(source, {**context, "source_folder": str(folder), "source_fingerprint": snapshot})
    data = (folder / "Yazad_Madan.docx").read_bytes()
    if hashlib.sha256(data).hexdigest() != snapshot["Yazad_Madan.docx"]:
        raise ValueError("The source resume changed during import. Run the import again.")
    return context, content, data, {"source_folder": str(folder), "source_fingerprint": snapshot, "audit_report": report}


def check_source(source, session):
    """Invalidate upload when the selected application or its tracker identity changes."""
    try:
        folder, paths, _ = application_files(source, session["source_folder"])
        if fingerprint(folder, paths) != session["source_fingerprint"]:
            raise ValueError("The source application changed. Import and review it again.")
        context = tracker_context(source, session["application_id"])
        expected = {key: session[key] for key in context}
        expected["url"] = session.get("tracker_url", session["url"])
        if context != expected:
            raise ValueError("The application tracker context changed. Import and review it again.")
    except (FileNotFoundError, OSError) as error:
        raise ValueError("The source application is unavailable. Import and review it again.") from error
