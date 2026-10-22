"""Run a fabricated local fixture without the private Resume workflow.

Author: Yazad Madan.
"""
import argparse
import hashlib
import json
import shutil
import tempfile
import threading
import uuid
from pathlib import Path

from docx import Document

from server import make_server

# Demo configuration
EXAMPLE_SOURCE = Path(__file__).resolve().parent / "examples" / "synthetic_source"
DEFAULT_PORT = 8767
RESUME_FILENAME = "Yazad_Madan.docx"
SESSION_FILENAME = "session.json"
DEMO_NOTICE = "SYNTHETIC DEMO ONLY. No real profile, audit, builder, or employer validation."


class DemoPipeline:
    def __init__(self, root):
        self.source = root / "source"
        shutil.copytree(EXAMPLE_SOURCE, self.source)
        self.data = root / "sessions"
        self.session_id = uuid.uuid4().hex
        self.session_folder = self.data / self.session_id
        self.session_folder.mkdir(parents=True)
        self.lock = threading.RLock()
        document = Document()
        document.core_properties.author = "Yazad Madan"
        document.core_properties.last_modified_by = "Yazad Madan"
        document.add_heading("Synthetic attachment", 0)
        document.add_paragraph(DEMO_NOTICE)
        document.save(self.session_folder / RESUME_FILENAME)
        self.session = {
            "id": self.session_id, "company": "Example Employer", "role": "Example Analyst",
            "mode": "sandbox", "state": "built", "upload_reviewed": True,
            "application_id": None, "fixture_notice": DEMO_NOTICE,
        }
        self.bind_origin(None)

    def bind_origin(self, origin):
        self.session["url"] = origin + "/fixture" if origin else ""
        (self.session_folder / SESSION_FILENAME).write_text(json.dumps(self.session), encoding="utf-8")

    def templates(self):
        return []

    def current(self):
        return dict(self.session)

    def folder(self, session_id):
        if session_id != self.session_id:
            raise ValueError("Unknown demo session.")
        return self.session_folder

    def read(self, session_id):
        self.folder(session_id)
        return self.current()

    def resume(self, session_id, *, for_upload=False):
        self.folder(session_id)
        if for_upload and self.session["url"] != getattr(self, "fixture_url", None):
            raise ValueError("Demo attachments are restricted to the local fixture.")
        return (self.session_folder / RESUME_FILENAME).read_bytes()

    def record_for(self, session_id):
        self.folder(session_id)
        raise ValueError("Demo sessions have no tracker or application record.")

    def create(self, body):
        raise ValueError("Demo uses one fixed synthetic session. Restart it to reset.")

    def import_application(self, body):
        raise ValueError("Demo cannot import real applications. Use the separate Resume workflow.")


def create_demo_server(root, port=DEFAULT_PORT):
    pipeline = DemoPipeline(root)
    server = make_server(pipeline, port)
    origin = f"http://127.0.0.1:{server.server_port}"
    pipeline.fixture_url = origin + "/fixture"
    pipeline.bind_origin(origin)
    pipeline.session["resume_sha256"] = hashlib.sha256(pipeline.resume(pipeline.session_id)).hexdigest()
    return server, pipeline


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--port", type=int, default=DEFAULT_PORT)
    args = parser.parse_args()
    with tempfile.TemporaryDirectory(prefix="portal-synthetic-demo-") as temporary:
        server, _ = create_demo_server(Path(temporary), args.port)
        print(DEMO_NOTICE, flush=True)
        print(f"Fixture: http://127.0.0.1:{server.server_port}/fixture", flush=True)
        print(f"Helper URL: http://127.0.0.1:{server.server_port}", flush=True)
        print(f"Pairing token: {server.token}", flush=True)
        print("Dummy attachment is preapproved only for this local fixture. Ctrl+C stops and resets the demo.", flush=True)
        try:
            server.serve_forever()
        except KeyboardInterrupt:
            pass
        finally:
            server.server_close()


if __name__ == "__main__":
    main()
