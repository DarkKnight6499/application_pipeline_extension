"""Preflight gates run on fabricated trackers, JD documents and field lists only."""
import hashlib
import shutil
import sys
import tempfile
import threading
import unittest
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

# Synthetic configuration
HERE = Path(__file__).resolve().parents[1]
AUTHOR = "Yazad Madan"
TRACKER_FILENAME = "Applications.xlsx"
OWN_ID = 801
OTHER_ID = 802
THIRD_ID = 803
OWN_COMPANY = "Synthetic Employer"
OWN_ROLE = "Synthetic Role"
OWN_LOCATION = "Newark, NJ"
OWN_LINK = "https://example.invalid/role/801"
OTHER_LINK = "https://example.invalid/role/802"
JD_FILENAME = "JD_Synthetic_Employer_Synthetic_Role.docx"
APPLICATION_ID_COLUMN = "Application ID"
COMPANY_COLUMN = "Company"
ROLE_COLUMN = "Role Title"
LINK_COLUMN = "Link"
LOCATION_COLUMN = "Location"
STATUS_COLUMN = "Status"
DATE_APPLIED_COLUMN = "Date Applied"
COLUMNS = [APPLICATION_ID_COLUMN, COMPANY_COLUMN, ROLE_COLUMN, LINK_COLUMN, LOCATION_COLUMN, STATUS_COLUMN, DATE_APPLIED_COLUMN]
REFERENCE_MODULES = ["identity_lib.py", "job_scout.py", "atomic_json.py", "archive_scrape_run.py", "keyword_matching.py"]

sys.path.insert(0, str(HERE))
sys.dont_write_bytecode = True
from test_support import INTEGRATION_SKIP_REASON, workflow_source
SOURCE = workflow_source()


def tracker_row(application_id, company, role, link, location, status, applied=None):
    return {APPLICATION_ID_COLUMN: application_id, COMPANY_COLUMN: company, ROLE_COLUMN: role, LINK_COLUMN: link,
            LOCATION_COLUMN: location, STATUS_COLUMN: status, DATE_APPLIED_COLUMN: applied}


OWN_ROW = tracker_row(OWN_ID, OWN_COMPANY, OWN_ROLE, OWN_LINK, OWN_LOCATION, "To Apply")


def field(label, intent=None, field_id=None, options=()):
    return {"id": field_id or label, "label": label, "type": "text", "intent": intent, "options": list(options)}


@unittest.skipIf(SOURCE is None, INTEGRATION_SKIP_REASON)
class PreflightTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.folder = self.root / "Applications" / "2026-10-04_Synthetic_Employer_Synthetic_Role"
        self.folder.mkdir(parents=True)
        self.reference = self.root / "_Reference"
        self.reference.mkdir()
        for name in REFERENCE_MODULES:
            shutil.copyfile(SOURCE / "_Reference" / name, self.reference / name)
        self.session = {"application_id": OWN_ID, "company": OWN_COMPANY, "role": OWN_ROLE, "tracker_url": OWN_LINK,
                        "url": OWN_LINK, "source_folder": str(self.folder)}
        import preflight
        self.preflight = preflight

    def tearDown(self):
        self.temporary.cleanup()

    def write_tracker(self, *rows):
        workbook = Workbook()
        workbook.properties.creator = AUTHOR
        workbook.properties.lastModifiedBy = AUTHOR
        workbook.active.append(COLUMNS)
        for row in rows:
            workbook.active.append([row[name] for name in COLUMNS])
        workbook.save(self.root / TRACKER_FILENAME)
        workbook.close()

    def write_jd(self, text):
        from docx import Document
        document = Document()
        document.core_properties.author = AUTHOR
        document.add_paragraph(text)
        document.save(self.folder / JD_FILENAME)

    def duplicate(self, *rows):
        self.write_tracker(OWN_ROW, *rows)
        return self.preflight.duplicate_gate(self.root, self.session, reference=self.reference)

    def test_link_match_blocks(self):
        items = self.duplicate(tracker_row(OTHER_ID, "Other Name", "Other Title", OWN_LINK, "Elsewhere", "Applied", date.today()))
        blocks = [item for item in items if item["severity"] == "block"]
        self.assertEqual(len(blocks), 1, items)
        self.assertTrue(blocks[0]["ack_required"])
        self.assertIn(str(OTHER_ID), blocks[0]["evidence"])

    def test_metadata_match_when_one_side_has_no_link_blocks(self):
        items = self.duplicate(tracker_row(OTHER_ID, OWN_COMPANY, OWN_ROLE, None, OWN_LOCATION, "Applied", date.today()))
        self.assertEqual([item["severity"] for item in items if item["gate"] == "duplicate"], ["block"], items)

    def test_different_links_same_title_not_duplicate(self):
        items = self.duplicate(tracker_row(OTHER_ID, OWN_COMPANY, OWN_ROLE, OTHER_LINK, OWN_LOCATION, "Rejected", date.today()))
        self.assertEqual([item for item in items if item["severity"] == "block"], [], items)

    def test_company_name_normalized_repeat_warns(self):
        self.session.update(company="Bank Of Synthetic", role="Analyst")
        self.write_tracker(tracker_row(OWN_ID, "Bank Of Synthetic", "Analyst", OWN_LINK, OWN_LOCATION, "To Apply"),
                           tracker_row(OTHER_ID, "BankOfSynthetic", "Different Role", OTHER_LINK, "Boston, MA", "Applied", date.today() - timedelta(days=30)),
                           tracker_row(THIRD_ID, "bank of synthetic", "Old Role", "https://example.invalid/role/803", "Boston, MA", "Applied",
                                       date.today() - timedelta(days=self.preflight.REPEAT_WINDOW_DAYS + 20)))
        items = self.preflight.duplicate_gate(self.root, self.session, reference=self.reference)
        warns = [item for item in items if item["severity"] == "warn"]
        self.assertEqual(len(warns), 1, items)
        self.assertIn(str(OTHER_ID), warns[0]["evidence"])
        self.assertNotIn(str(THIRD_ID), warns[0]["evidence"])
        self.assertEqual([item for item in items if item["severity"] == "block"], [])

    def test_blocked_status_blocks(self):
        for status in ("Blocked", "Skipped", "Expired"):
            with self.subTest(status=status):
                self.write_tracker({**OWN_ROW, STATUS_COLUMN: status})
                items = self.preflight.duplicate_gate(self.root, self.session, reference=self.reference)
                blocks = [item for item in items if item["gate"] == "tracker_status"]
                self.assertEqual([item["severity"] for item in blocks], ["block"], items)

    def test_no_sponsorship_jd_blocks_with_evidence(self):
        self.write_jd("Responsibilities include reporting. We do not sponsor visas for this position.")
        items = self.preflight.sponsorship_gate(self.root, self.session, reference=self.reference)
        self.assertEqual([item["severity"] for item in items], ["block"], items)
        self.assertIn("do not sponsor", items[0]["evidence"])
        self.assertTrue(items[0]["ack_required"])

    def test_sponsorship_clear_jd_has_no_block(self):
        self.write_jd("Responsibilities include reporting and analysis.")
        items = self.preflight.sponsorship_gate(self.root, self.session, reference=self.reference)
        self.assertEqual([item for item in items if item["severity"] == "block"], [])

    def test_sponsorship_without_jd_says_not_checked(self):
        items = self.preflight.sponsorship_gate(self.root, self.session, reference=self.reference)
        self.assertEqual([item["severity"] for item in items], ["warn"])
        self.assertFalse(items[0]["ack_required"])

    def test_jd_outside_applications_is_refused(self):
        self.session["source_folder"] = str(self.root)
        items = self.preflight.sponsorship_gate(self.root, self.session, reference=self.reference)
        self.assertEqual([item["severity"] for item in items], ["warn"])

    def test_knockout_years_and_clearance_warn(self):
        items = self.preflight.knockout_gate([
            field("Do you have 8+ years of experience in treasury?", field_id="years"),
            field("Do you have 3 years of experience in treasury?", field_id="few"),
            field("Do you hold an active security clearance?", field_id="clearance"),
            field("Do you have a PhD in finance?", field_id="phd"),
            field("Must be a U.S. citizen only", field_id="citizen"),
            field("Is $150,000 your minimum salary?", field_id="salary"),
            field("Are you able to work on-site five days a week?", field_id="onsite"),
            field("First name", field_id="name")])
        self.assertCountEqual([item["field_id"] for item in items], ["years", "clearance", "phd", "citizen", "salary", "onsite"])
        for item in items:
            self.assertEqual(item["severity"], "warn")
            self.assertTrue(item["ack_required"])
            self.assertEqual(item["gate"], "knockout")

    def test_status_question_forced_manual(self):
        items = self.preflight.status_gate([field("Are you a US citizen or green card holder?", intent="status_question", field_id="status"),
                                            field("First name", field_id="name")])
        self.assertEqual([(item["field_id"], item["severity"], item["force_manual"]) for item in items], [("status", "warn", True)])
        self.write_tracker(OWN_ROW)
        result = self.preflight.run_preflight(self.root, self.session, [field("Are you a citizen?", intent="status_question", field_id="status")],
                                              reference=self.reference)
        self.assertIn("status", result["manual_field_ids"])

    def test_salary_history_forced_manual(self):
        labels = {"history": "What was your salary history?", "current": "Current salary", "dob": "Date of birth", "age": "What is your age?",
                  "ssn": "Social Security number", "license": "Driver's license number", "criminal": "Have you been convicted of a crime? Criminal history"}
        items = self.preflight.sensitive_gate([field(label, field_id=key) for key, label in labels.items()] + [field("Percentage of time", field_id="ok")])
        self.assertCountEqual([item["field_id"] for item in items], list(labels))
        for item in items:
            self.assertEqual(item["severity"], "warn")
            self.assertTrue(item["force_manual"])
            self.assertEqual(item["message"], "Answer this yourself.")
        self.write_tracker(OWN_ROW)
        result = self.preflight.run_preflight(self.root, self.session, [field("Salary history", field_id="history")], reference=self.reference)
        self.assertIn("history", result["manual_field_ids"])

    def test_run_preflight_without_fields_skips_field_gates(self):
        self.write_tracker(OWN_ROW)
        self.write_jd("Plain description.")
        result = self.preflight.run_preflight(self.root, self.session, None, reference=self.reference)
        self.assertEqual(result["items"], [])
        self.assertEqual(result["manual_field_ids"], [])
        self.assertFalse(result["ack_required"])

    def test_missing_tracker_is_reported_not_silent(self):
        result = self.preflight.run_preflight(self.root, self.session, None, reference=self.reference)
        gates = {entry["gate"]: entry for entry in result["items"]}
        self.assertEqual(gates["duplicate"]["severity"], "warn")
        self.assertIn("not checked", gates["duplicate"]["message"])
        self.assertFalse(result["ack_required"])

    def test_preflight_never_writes_tracker(self):
        self.write_tracker(OWN_ROW, tracker_row(OTHER_ID, OWN_COMPANY, OWN_ROLE, OWN_LINK, OWN_LOCATION, "Applied", date.today()))
        self.write_jd("We do not sponsor visas.")
        path = self.root / TRACKER_FILENAME
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        stamp = path.stat().st_mtime_ns
        result = self.preflight.run_preflight(self.root, self.session, [field("Years of experience 9+ years")], reference=self.reference)
        self.assertTrue(result["ack_required"])
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)
        self.assertEqual(path.stat().st_mtime_ns, stamp)
        self.assertEqual(sorted(entry.name for entry in self.root.iterdir()), sorted([TRACKER_FILENAME, "Applications", "_Reference"]))

    def test_read_only_workbook_calls(self):
        source = (HERE / "preflight.py").read_text(encoding="utf-8")
        calls = [line for line in source.splitlines() if "load_workbook(" in line and "import" not in line]
        self.assertTrue(calls)
        self.assertTrue(all("read_only=True" in line for line in calls), calls)


@unittest.skipIf(SOURCE is None, INTEGRATION_SKIP_REASON)
class PreflightPanelTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from playwright.sync_api import sync_playwright
        from server import Pipeline, make_server
        cls.temp = tempfile.TemporaryDirectory()
        cls.pipeline = Pipeline(SOURCE, Path(cls.temp.name) / "output")
        cls.server = make_server(cls.pipeline, 0, "preflight-test-token")
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.url = f"http://127.0.0.1:{cls.server.server_port}"
        cls.pipeline.create({"company": "Synthetic Employer", "role": "Treasury Analyst", "jd": "Treasury analyst: Python and SQL.",
                             "template": "Resume_Draft_ALM_Treasury_Liquidity.json", "url": cls.url + "/fixture"})
        cls.playwright = sync_playwright().start()
        import os
        cls.browser = cls.playwright.chromium.launch(headless=True, channel=os.environ.get("PORTAL_TEST_BROWSER") or None)

    @classmethod
    def tearDownClass(cls):
        cls.browser.close()
        cls.playwright.stop()
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join()
        cls.temp.cleanup()

    def setUp(self):
        self.context = self.browser.new_context(viewport={"width": 1440, "height": 1050})
        self.page = self.context.new_page()
        self.errors = []
        self.page.on("pageerror", lambda error: self.errors.append(str(error)))
        self.page.goto(self.url + "/fixture")

    def tearDown(self):
        self.context.close()
        self.assertEqual(self.errors, [])

    def test_fill_disabled_until_ack(self):
        self.page.evaluate("""() => {
          const label = document.createElement('label'); label.textContent = 'Do you have 9+ years of experience in treasury?';
          const input = document.createElement('input'); label.append(input); document.querySelector('form').append(label);
        }""")
        self.page.locator("#scan").click()
        host = self.page.locator("#portal-panel-host")
        fill = host.get_by_role("button", name="Fill selected fields")
        fill.wait_for()
        ack = host.get_by_role("checkbox", name="Acknowledge preflight item", exact=False).first
        ack.wait_for()
        self.assertTrue(fill.is_disabled())
        ack.check()
        self.assertFalse(fill.is_disabled())
        ack.uncheck()
        self.assertTrue(fill.is_disabled())
