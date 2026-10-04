"""Read-only preflight gates that run before any value reaches a portal; ideas from career-ops-hq/career-ops (MIT), written fresh."""
from __future__ import annotations

import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

sys.dont_write_bytecode = True

# Tracker and gate configuration
COMPANY_COLUMN = "Company"
ROLE_COLUMN = "Role Title"
LINK_COLUMN = "Link"
LOCATION_COLUMN = "Location"
STATUS_COLUMN = "Status"
APPLICATION_ID_COLUMN = "Application ID"
DATE_APPLIED_COLUMN = "Date Applied"
REQUIRED_COLUMNS = [APPLICATION_ID_COLUMN, COMPANY_COLUMN, ROLE_COLUMN, STATUS_COLUMN]
TRACKER_FILENAME = "Applications.xlsx"
REPEAT_WINDOW_DAYS = 180
YEARS_THRESHOLD = 7
BLOCKED_TRACKER_STATUSES = {"blocked", "skipped", "expired"}
REPEAT_STATUSES = {"applied", "screening", "to apply"}
DATED_REPEAT_STATUSES = {"applied", "screening"}
SENSITIVE_MESSAGE = "Answer this yourself."
MAX_EVIDENCE = 300


def item(gate, severity, message, evidence="", ack_required=False, **extra):
    return {"gate": gate, "severity": severity, "message": message, "evidence": str(evidence)[:MAX_EVIDENCE],
            "ack_required": ack_required, **extra}


def _reference_dir(source, reference):
    return Path(reference) if reference else Path(source) / "_Reference"


def _import(reference, name):
    import importlib
    if str(reference) not in sys.path:
        sys.path.insert(0, str(reference))
    return importlib.import_module(name)


def _normal(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _read_tracker(source):
    from openpyxl import load_workbook
    workbook = load_workbook(Path(source) / TRACKER_FILENAME, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        headers = list(next(rows, ()))
        if any(headers.count(column) != 1 for column in REQUIRED_COLUMNS):
            raise ValueError("The tracker is missing a required column or repeats one.")
        return [dict(zip(headers, values)) for values in rows if any(value is not None for value in values)]
    finally:
        workbook.close()


def _as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip()[:10])
    except ValueError:
        return None


def duplicate_gate(source, session, *, reference=None):
    identity = _import(_reference_dir(source, reference), "identity_lib")
    rows = _read_tracker(source)
    own_id = session.get("application_id")
    own = next((row for row in rows if own_id is not None and row.get(APPLICATION_ID_COLUMN) == own_id), {})
    others = [row for row in rows if own_id is None or row.get(APPLICATION_ID_COLUMN) != own_id]
    items = []
    status = str(own.get(STATUS_COLUMN) or "").strip()
    if status.lower() in BLOCKED_TRACKER_STATUSES:
        items.append(item("tracker_status", "block", f"The tracker marks this application {status}. Confirm before filling anything.",
                          f"Application ID {own_id}, Status {status}", True))
    company = own.get(COMPANY_COLUMN) or session.get("company")
    role = own.get(ROLE_COLUMN) or session.get("role")
    location = own.get(LOCATION_COLUMN) or ""
    link = own.get(LINK_COLUMN) or session.get("tracker_url") or session.get("url") or ""
    index = identity.IdentityIndex()
    for row in others:
        index.add(row.get(COMPANY_COLUMN), row.get(ROLE_COLUMN), row.get(LOCATION_COLUMN), row.get(LINK_COLUMN), row.get(APPLICATION_ID_COLUMN))
    match = index.lookup(company, role, location, link)
    if match is not None:
        items.append(item("duplicate", "block", "Another tracker row identifies the same posting. Confirm this is not a repeat application.",
                          f"Application ID {match}", True))
    cutoff = date.today() - timedelta(days=REPEAT_WINDOW_DAYS)
    repeats = []
    for row in others:
        row_status = str(row.get(STATUS_COLUMN) or "").strip().lower()
        if _normal(row.get(COMPANY_COLUMN)) != _normal(company) or row_status not in REPEAT_STATUSES or row.get(APPLICATION_ID_COLUMN) == match:
            continue
        applied = _as_date(row.get(DATE_APPLIED_COLUMN))
        if row_status in DATED_REPEAT_STATUSES and applied is not None and applied < cutoff:
            continue
        repeats.append(str(row.get(APPLICATION_ID_COLUMN)))
    if repeats:
        items.append(item("repeat_company", "warn", f"Other applications to this company are open or recent (last {REPEAT_WINDOW_DAYS} days).",
                          "Application IDs " + ", ".join(repeats)))
    return items


def _jd_text(source, session):
    folder = session.get("source_folder")
    if not folder:
        return None
    applications = (Path(source) / "Applications").resolve()
    folder = Path(folder).resolve()
    if not folder.is_relative_to(applications) or folder == applications or not folder.is_dir():
        return None
    documents = list(folder.glob("JD_*.docx"))
    if len(documents) != 1:
        return None
    from docx import Document
    document = Document(documents[0])
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        parts += [cell.text for row in table.rows for cell in row.cells]
    return "\n".join(parts)


def sponsorship_gate(source, session, *, reference=None):
    text = _jd_text(source, session)
    if not text:
        return [item("sponsorship", "warn", "The JD text was not available, so its sponsorship language was not checked.")]
    scout = _import(_reference_dir(source, reference), "job_scout")
    verdict, evidence = scout.classify_sponsorship(text)
    if verdict == "Blocked":
        return [item("sponsorship", "block", "The JD states a no-sponsorship or citizenship requirement. Confirm before filling.", evidence, True)]
    return []


def _text(entry):
    return str(entry.get("label") or "")


def _norm_label(entry):
    return re.sub(r"\s+", " ", _text(entry).lower().replace("’", "'")).strip()


KNOCKOUTS = [
    ("Asks for a PhD or doctorate.", re.compile(r"\b(ph\.?\s?d|doctorate|doctoral)\b")),
    ("Asks about a security clearance.", re.compile(r"\b(security clearance|clearance|top secret|ts/sci)\b")),
    ("Looks restricted to citizens or green card holders.", re.compile(
        r"(u\.?s\.?\s+citizens?\s+(only|required)|must be a (u\.?s\.?\s+)?citizen|citizens? only|green card (holders? )?(only|required)|permanent residents? (only|required))")),
    ("Asks for an on-site or in-office commitment.", re.compile(
        r"\b(on-?\s?site|in[- ]office|in-person)\b.*\b(required|must|require|able|willing|days|week)\b|\b(required|must|require|able|willing|days|week)\b.*\b(on-?\s?site|in[- ]office)\b")),
]
YEARS = re.compile(r"(\d+)\s*\+?\s*(?:or more\s+|plus\s+)?years?")
SALARY = re.compile(r"\b(salary|compensation|pay)\b")
MONEY = re.compile(r"\$\s?\d|\d[\d,]*\s?k\b|\b\d{2,3}[,.]\d{3}\b")


def knockout_gate(fields):
    items = []
    for entry in fields or []:
        label = _norm_label(entry)
        reasons = []
        years = [int(number) for number in YEARS.findall(label)]
        if "experience" in label and any(number >= YEARS_THRESHOLD for number in years):
            reasons.append(f"Asks for {max(years)} or more years of experience.")
        for reason, pattern in KNOCKOUTS:
            if pattern.search(label):
                reasons.append(reason)
        if SALARY.search(label) and MONEY.search(label):
            reasons.append("States a salary figure.")
        for reason in reasons:
            items.append(item("knockout", "warn", reason + " Acknowledge before filling.", _text(entry), True, field_id=entry.get("id")))
    return items


def status_gate(fields):
    return [item("status_question", "warn", "This asks about status or citizenship. Answer this yourself.", _text(entry), False,
                 field_id=entry.get("id"), force_manual=True)
            for entry in fields or [] if entry.get("intent") == "status_question"]


SENSITIVE = re.compile(
    r"(salary history|(current|present|previous|prior|last)\s+(salary|compensation|pay)|date of birth|birth\s?date|\bdob\b|birthday|\bage\b"
    r"|social security|\bssn\b|driver'?s?\s+licen[sc]e|criminal|convicted|felony|arrest)")


def sensitive_gate(fields):
    return [item("sensitive", "warn", SENSITIVE_MESSAGE, _text(entry), False, field_id=entry.get("id"), force_manual=True)
            for entry in fields or [] if SENSITIVE.search(_norm_label(entry))]


def run_preflight(source, session, fields, *, reference=None):
    items = []
    for name, gate in (("duplicate", duplicate_gate), ("sponsorship", sponsorship_gate)):
        try:
            items += gate(source, session, reference=reference)
        except (OSError, ImportError, ValueError, KeyError) as error:
            items.append(item(name, "warn", f"The {name} check could not run, so it was not checked.", error))
    if fields is not None:
        items += knockout_gate(fields) + status_gate(fields) + sensitive_gate(fields)
    manual = sorted({entry["field_id"] for entry in items if entry.get("force_manual") and entry.get("field_id") is not None})
    return {"items": items, "manual_field_ids": manual, "ack_required": any(entry["ack_required"] for entry in items)}
